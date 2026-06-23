from dataclasses import dataclass
from pathlib import Path

import aiofiles
import openpyxl
import pypdf
from docx import Document as DocxDocument
from fastapi import UploadFile

from knowledge_common.exceptions.exception import ServiceException
from knowledge_common.utils.log_util import logger
import tempfile

@dataclass
class TempFileResult:
    """UploadFile 写入本地临时文件的结果"""

    path: str  # 临时文件路径
    size: int  # 文件大小（字节数）


class FileUtil:
    @staticmethod
    async def save_upload_to_temp(
        file: UploadFile,
        max_size: int,
        chunk_size: int = 65536,
        temp_dir: str | None = None,
    ) -> TempFileResult:
        """
        将 UploadFile 分块写入本地临时文件，避免大文件全部加载到内存。

        调用方需自行在 finally 块中清理临时文件。

        :param file: FastAPI UploadFile
        :param max_size: 文件大小上限（字节），需由业务调用方指定
        :param chunk_size: 每块读取大小（字节），默认 64KB
        :param temp_dir: 临时文件存放目录，默认使用系统临时目录
        :return: TempFileResult
        :raises ServiceException: 文件为空或大小超限时抛出
        """


        with tempfile.NamedTemporaryFile(delete=False, dir=temp_dir) as tmp_file:
            temp_path = tmp_file.name
            total_size = 0
            while True:
                chunk = await file.read(chunk_size)
                if not chunk:
                    break
                total_size += len(chunk)
                if total_size > max_size:
                    raise ServiceException(f'文件大小超过{max_size // (1024 * 1024)}MB限制')
                tmp_file.write(chunk)

        if total_size == 0:
            raise ServiceException('上传文件内容为空')

        return TempFileResult(path=temp_path, size=total_size)

    @staticmethod
    async def file_chunk_stream(file_path: Path, chunk_size: int = 256 * 1024):
        """异步流式读取文件，内存中仅保留一个 chunk 大小的缓冲区"""
        async with aiofiles.open(file_path, 'rb') as f:
            while chunk := await f.read(chunk_size):
                yield chunk

    @staticmethod
    async def clean_temp_file(temp_path: str):
        """删除临时文件"""
        if Path(temp_path).exists():
            Path(temp_path).unlink()

    @staticmethod
    def resolve_total_pages(file_path: str, ext: str) -> int:
        """
        解析文件总页数，供 MinerU 解析任务按页分段参考（MAX_PAGES_PER_SEGMENT）。

        各格式处理策略：
        - pdf: 有真实页数概念，通过 pypdf.PdfReader 读取实际页数。
        - docx: 页数依赖渲染引擎（字体、边距、纸张大小 等），无法精确计算；
                此处统计文档 XML 中显式分页符 <w:br w:type="page"/> 数量 + 1，
                作为最小页数估算，未出现显式分页符时返回 1。
        - xlsx: 没有传统意义上的"页"概念，以工作表数量作为近似页数。
        - doc: 旧版 Word 97-2003 二进制格式，python-docx 无法读取，
                暂无可跨平台使用的精确页数解析库，直接返回 1。
        - 其他格式: 默认返回 1，避免影响整体流程。

        :param file_path: 本地文件路径
        :param ext: 文件后缀（小写，如 'pdf'、'docx'）
        :return: 总页数，解析失败或不支持格式时返回 1
        """
        if ext == 'pdf':
            try:
                reader = pypdf.PdfReader(file_path)
                return len(reader.pages)  # 返回 PDF 实际页数
            except Exception as e:
                logger.warning(f'PDF 页数解析失败: {e}')
                return 1  # 解析失败时兜底返回 1，不影响后续流程

        if ext == 'docx':
            try:
                doc = DocxDocument(file_path)
                # 显式分页符 <w:br w:type="page"/> 出现次数 + 首页 = 最小页数
                page_breaks = doc.element.xml.count('w:type="page"')
                return page_breaks + 1
            except Exception as e:
                logger.warning(f'DOCX 页数解析失败: {e}')
                return 1  # 解析失败时兜底返回 1，不影响后续流程

        if ext == 'xlsx':
            try:
                wb = openpyxl.load_workbook(file_path, read_only=True)
                count = len(wb.sheetnames)  # 工作表数量作为近似页数
                wb.close()
                return count
            except Exception as e:
                logger.warning(f'XLSX 页数解析失败: {e}')
                return 1  # 解析失败时兜底返回 1，不影响后续流程

        return 1  # 未支持格式默认返回 1